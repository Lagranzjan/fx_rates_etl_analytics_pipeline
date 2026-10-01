"""
Pipeline ETL: kursy walut NBP (tabela A) -> SQLite -> analityka SQL.

Przykłady użycia:
    python nbp_pipeline.py                              # USD, EUR, CHF, GBP, ostatnie 60 dni
    python nbp_pipeline.py -c USD EUR -d 120            # wybrane waluty i zakres
    python nbp_pipeline.py --xlsx kursy.xlsx --plot     # Excel (openpyxl) + wykres PNG (matplotlib)
    python nbp_pipeline.py --csv wyniki.csv             # CSV z przecinkiem dziesiętnym (polski Excel)

Kolejne uruchomienia dociągają tylko brakujące dni (ładowanie przyrostowe).
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import sqlite3
import time
import urllib.error
import urllib.request
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

log = logging.getLogger("nbp")

API_BASE = "https://api.nbp.pl/api/exchangerates/rates/a"
MAX_RANGE_DAYS = 360          # limit API NBP dla zakresu dat to 367 dni
TIMEOUT_S = 10
RETRIES = 3
RATE_SCALE = 10_000           # kurs NBP ma 4 miejsca po przecinku -> przechowujemy jako INTEGER


# --------------------------------------------------------------------------
# 1. EXTRACT
# --------------------------------------------------------------------------
def fetch_json(url: str) -> dict | None:
    """Pobiera JSON z ponawianiem prób. Zwraca None, gdy API odpowiada 404 (brak danych w zakresie)."""
    req = urllib.request.Request(url, headers={"User-Agent": "nbp-etl/1.0", "Accept": "application/json"})
    last_err: Exception | None = None
    for attempt in range(1, RETRIES + 1):
        try:
            with urllib.request.urlopen(req, timeout=TIMEOUT_S) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 404:          # NBP zwraca 404, gdy w zakresie nie ma notowań (weekend/święta)
                return None
            last_err = e
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as e:
            last_err = e
        log.warning("Próba %d/%d nieudana (%s): %s", attempt, RETRIES, url, last_err)
        time.sleep(2 ** (attempt - 1))
    raise RuntimeError(f"Nie udało się pobrać danych z {url}") from last_err


def date_chunks(start: date, end: date, size: int = MAX_RANGE_DAYS):
    """Dzieli zakres dat na kawałki mieszczące się w limicie API."""
    cur = start
    while cur <= end:
        chunk_end = min(cur + timedelta(days=size - 1), end)
        yield cur, chunk_end
        cur = chunk_end + timedelta(days=1)


def extract(currency: str, start: date, end: date) -> list[dict]:
    """Pobiera surowe notowania waluty w podanym zakresie dat."""
    rows: list[dict] = []
    for s, e in date_chunks(start, end):
        url = f"{API_BASE}/{currency.lower()}/{s.isoformat()}/{e.isoformat()}/?format=json"
        payload = fetch_json(url)
        if payload:
            code = payload.get("code", currency.upper())
            rows.extend({"code": code, **r} for r in payload.get("rates", []))
    log.info("%s: pobrano %d surowych rekordów (%s – %s)", currency.upper(), len(rows), start, end)
    return rows


# --------------------------------------------------------------------------
# 2. TRANSFORM
# --------------------------------------------------------------------------
def transform(raw_rows: list[dict]) -> list[tuple[str, str, int]]:
    """Czyści i waliduje dane. Zwraca krotki (currency, rate_date, rate_e4)."""
    clean: list[tuple[str, str, int]] = []
    seen: set[tuple[str, str]] = set()
    rejected = 0

    for r in raw_rows:
        try:
            code = str(r["code"]).upper()
            rate_date = date.fromisoformat(r["effectiveDate"]).isoformat()
            mid = Decimal(str(r["mid"]))
        except (KeyError, ValueError, TypeError, InvalidOperation):
            rejected += 1
            continue

        if not mid.is_finite() or mid <= 0 or (code, rate_date) in seen:
            rejected += 1
            continue

        seen.add((code, rate_date))
        clean.append((code, rate_date, int((mid * RATE_SCALE).to_integral_value())))

    if rejected:
        log.warning("Odrzucono %d rekordów (braki, błędne wartości lub duplikaty)", rejected)
    return clean


# --------------------------------------------------------------------------
# 3. LOAD
# --------------------------------------------------------------------------
SCHEMA = """
CREATE TABLE IF NOT EXISTS currency_rates (
    currency  TEXT    NOT NULL,
    rate_date TEXT    NOT NULL,              -- ISO 8601: YYYY-MM-DD
    rate_e4   INTEGER NOT NULL CHECK (rate_e4 > 0),   -- kurs * 10000 (bez błędów float)
    PRIMARY KEY (currency, rate_date)
) WITHOUT ROWID;
"""


def init_db(conn: sqlite3.Connection) -> None:
    conn.executescript(SCHEMA)


def last_loaded_date(conn: sqlite3.Connection, currency: str) -> date | None:
    row = conn.execute(
        "SELECT MAX(rate_date) FROM currency_rates WHERE currency = ?", (currency.upper(),)
    ).fetchone()
    return date.fromisoformat(row[0]) if row and row[0] else None


def load(conn: sqlite3.Connection, records: list[tuple[str, str, int]]) -> int:
    """Upsert – ponowne uruchomienie nie duplikuje ani nie kasuje historii."""
    with conn:  # transakcja: commit albo rollback
        conn.executemany(
            """
            INSERT INTO currency_rates (currency, rate_date, rate_e4)
            VALUES (?, ?, ?)
            ON CONFLICT (currency, rate_date) DO UPDATE SET rate_e4 = excluded.rate_e4;
            """,
            records,
        )
    return len(records)


# --------------------------------------------------------------------------
# 4. ANALYTICS
# --------------------------------------------------------------------------
ANALYTICS_SQL = """
WITH base AS (
    SELECT currency, rate_date, rate_e4 / 10000.0 AS rate
    FROM currency_rates
),
calc AS (
    SELECT
        currency,
        rate_date,
        rate,
        LAG(rate, 1) OVER (PARTITION BY currency ORDER BY rate_date) AS prev_rate,
        LAG(rate, 7) OVER (PARTITION BY currency ORDER BY rate_date) AS rate_7d_ago,
        -- średnia krocząca tylko dla pełnego 7-elementowego okna, inaczej NULL
        CASE WHEN COUNT(*) OVER w7 = 7 THEN AVG(rate) OVER w7 END AS ma7,
        CASE WHEN COUNT(*) OVER w7 = 7 THEN MIN(rate) OVER w7 END AS min7,
        CASE WHEN COUNT(*) OVER w7 = 7 THEN MAX(rate) OVER w7 END AS max7,
        ROW_NUMBER() OVER (PARTITION BY currency ORDER BY rate_date DESC) AS rn
    FROM base
    WINDOW w7 AS (PARTITION BY currency ORDER BY rate_date
                  ROWS BETWEEN 6 PRECEDING AND CURRENT ROW)
)
SELECT
    currency,
    rate_date,
    rate,
    prev_rate,
    ROUND((rate - prev_rate) / prev_rate * 100, 2)        AS daily_change_pct,
    ROUND(ma7, 4)                                          AS ma7,
    ROUND(min7, 4)                                         AS min7,
    ROUND(max7, 4)                                         AS max7,
    ROUND((rate - rate_7d_ago) / rate_7d_ago * 100, 2)     AS trend_7d_pct
FROM calc
WHERE rn <= :limit
ORDER BY currency, rate_date DESC;
"""

COLUMNS = ["currency", "rate_date", "rate", "prev_rate", "daily_change_pct",
           "ma7", "min7", "max7", "trend_7d_pct"]


def analyze(conn: sqlite3.Connection, limit: int = 10) -> list[tuple]:
    return conn.execute(ANALYTICS_SQL, {"limit": limit}).fetchall()


def fmt(value, spec: str = ".4f", signed: bool = False, suffix: str = "") -> str:
    """Bezpieczne formatowanie – NULL (None) zamienia na 'N/A'."""
    if value is None:
        return "N/A"
    return f"{value:{'+' if signed else ''}{spec}}{suffix}"


def print_report(rows: list[tuple]) -> None:
    header = (f"{'Waluta':<6} | {'Data':<10} | {'Kurs':>8} | {'Poprz.':>8} | {'Zmiana':>8} | "
              f"{'Śr.7d':>8} | {'Min7':>8} | {'Max7':>8} | {'Trend7d':>8}")
    print("\n--- WYNIKI ANALIZY SQL ---")
    print(header)
    print("-" * len(header))
    for cur, d, rate, prev, chg, ma7, mn, mx, tr in rows:
        print(f"{cur:<6} | {d:<10} | {fmt(rate):>8} | {fmt(prev):>8} | "
              f"{fmt(chg, '.2f', True, '%'):>8} | {fmt(ma7):>8} | {fmt(mn):>8} | "
              f"{fmt(mx):>8} | {fmt(tr, '.2f', True, '%'):>8}")


# --------------------------------------------------------------------------
# Eksport i wizualizacja (opcjonalne)
# --------------------------------------------------------------------------
def export_csv(rows: list[tuple], path: str) -> None:
    # utf-8-sig (BOM) + średnik: polski Excel poprawnie czyta znaki i kolumny
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(COLUMNS)
        # przecinek dziesiętny: polski Excel kropkę w "4.6182" czytałby jako datę
        writer.writerows(
            [str(v).replace(".", ",") if isinstance(v, float) else v for v in row]
            for row in rows
        )
    log.info("Zapisano CSV: %s (%d wierszy)", path, len(rows))


def export_xlsx(rows: list[tuple], path: str) -> None:
    """Zapisuje skoroszyt Excela: zakładka na walutę + natywny wykres Excela."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        log.warning("Brak openpyxl – pomijam XLSX (pip install openpyxl)")
        return

    headers = ["Data", "Kurs", "Poprzedni kurs", "Zmiana dzienna %", "Średnia 7d",
               "Min 7d", "Max 7d", "Trend 7d %"]
    by_cur: dict[str, list[tuple]] = {}
    for r in rows:
        by_cur.setdefault(r[0], []).append(r)

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F4E78")

    for cur, data in by_cur.items():
        data.sort(key=lambda r: r[1])                      # rosnąco po dacie
        ws = wb.create_sheet(title=f"{cur}")
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")

        for _, d, rate, prev, chg, ma7, mn, mx, tr in data:
            ws.append([date.fromisoformat(d), rate, prev, chg, ma7, mn, mx, tr])

        n = len(data)
        for row in ws.iter_rows(min_row=2, max_row=n + 1):
            row[0].number_format = "yyyy-mm-dd"
            for i in (1, 2, 4, 5, 6):
                row[i].number_format = "0.0000"
            for i in (3, 7):
                row[i].number_format = '+0.00"%";-0.00"%";0.00"%"'

        for col, w in zip("ABCDEFGH", (12, 10, 16, 18, 12, 10, 10, 12)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:H{n + 1}"

        chart = LineChart()
        chart.title = f"{cur}/PLN"
        chart.height, chart.width = 8, 18
        chart.y_axis.title = "PLN"
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
        chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=n + 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        chart.x_axis.number_format = "mm-dd"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        ws.add_chart(chart, "J2")

    wb.save(path)
    log.info("Zapisano Excel: %s (%d zakładek)", path, len(by_cur))


def plot(rows: list[tuple], path: str = "nbp_chart.png") -> None:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        log.warning("Brak matplotlib – pomijam wykres (pip install matplotlib)")
        return

    by_cur: dict[str, list[tuple]] = {}
    for r in rows:
        by_cur.setdefault(r[0], []).append(r)

    fig, axes = plt.subplots(len(by_cur), 1, figsize=(10, 3.2 * len(by_cur)), squeeze=False)
    for ax, (cur, data) in zip(axes[:, 0], by_cur.items()):
        data.sort(key=lambda r: r[1])
        dates = [date.fromisoformat(r[1]) for r in data]
        ax.plot(dates, [r[2] for r in data], label=f"{cur}/PLN")
        ma = [(d, r[5]) for d, r in zip(dates, data) if r[5] is not None]
        if ma:
            ax.plot(*zip(*ma), linestyle="--", label="Średnia 7d")
        ax.set_title(f"{cur}/PLN")
        ax.grid(alpha=0.3)
        ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    plt.close(fig)
    log.info("Zapisano wykres: %s", path)


# --------------------------------------------------------------------------
# Orkiestracja
# --------------------------------------------------------------------------
def run(currencies: list[str], days: int, db_path: str) -> sqlite3.Connection:
    today = date.today()
    conn = sqlite3.connect(db_path)
    init_db(conn)

    for cur in currencies:
        last = last_loaded_date(conn, cur)
        start = last + timedelta(days=1) if last else today - timedelta(days=days)
        if start > today:
            log.info("%s: dane aktualne (ostatnia data: %s)", cur.upper(), last)
            continue
        try:
            raw = extract(cur, start, today)
        except RuntimeError as e:
            log.error("%s: pomijam – %s", cur.upper(), e)
            continue
        clean = transform(raw)
        log.info("%s: załadowano %d rekordów", cur.upper(), load(conn, clean))
    return conn


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="ETL kursów walut NBP -> SQLite")
    p.add_argument("-c", "--currencies", nargs="+", default=["USD", "EUR", "CHF", "GBP"],
                   help="kody walut (domyślnie: USD EUR CHF GBP)")
    p.add_argument("-d", "--days", type=int, default=60,
                   help="ile dni wstecz przy pierwszym uruchomieniu (domyślnie 60)")
    p.add_argument("--db", default="nbp_analytics.db", help="ścieżka do pliku SQLite")
    p.add_argument("--limit", type=int, default=10, help="ile ostatnich dni w raporcie na walutę")
    p.add_argument("--csv", metavar="PLIK", help="eksportuj pełną analitykę do CSV")
    p.add_argument("--xlsx", metavar="PLIK", help="eksportuj do Excela (.xlsx, zakładka na walutę + wykresy)")
    p.add_argument("--plot", action="store_true", help="zapisz wykres nbp_chart.png")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s %(message)s")

    conn = run(args.currencies, args.days, args.db)
    try:
        print_report(analyze(conn, args.limit))
        if args.csv or args.xlsx or args.plot:
            full = analyze(conn, limit=10**9)
            if args.csv:
                export_csv(full, args.csv)
            if args.xlsx:
                export_xlsx(full, args.xlsx)
            if args.plot:
                plot(full)
    finally:
        conn.close()


if __name__ == "__main__":
    main()